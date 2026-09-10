"""纯检测报告与中心 Excel 三表更新。"""

from __future__ import annotations

import copy
import json
import os
from datetime import datetime
from pathlib import Path
from typing import Any

from pipeline.core.errors import ExcelPendingError, PipelineError
from pipeline.core.io_utils import atomic_write_text
from pipeline.core.stage_runner import StageOutcome
from pipeline.steps.reporting import (
    EXCEL_DURATION_FORMAT,
    _acquire_lock,
    _normalize_excel_durations,
    _normalize_excel_metrics,
    validate_excel_template,
)

from .data_registry import _collect_registered_statistics, _load_registry


SPLITS = ("train", "val")
CLASS_NAMES = {0: "paper", 1: "liquid", 2: "metal"}


def _stage(context: Any, name: str) -> Any:
    """读取已成功的阶段记录。"""
    stage = context.state.get_stage(name)
    if stage.status.value not in {"succeeded", "skipped"}:
        raise PipelineError(f"报告需要已完成阶段:{name}")
    return stage


def _metric(metrics: dict[str, Any], *keys: str) -> Any:
    """兼容不同 Ultralytics 版本的指标键名。"""
    return next((metrics[key] for key in keys if metrics.get(key) is not None), None)


def _path(context: Any, value: Any) -> str:
    """将运行路径映射为用户可访问的 UNC 展示路径。"""
    return "" if value in (None, "") else context.mapper.map_path(str(value))


def _invocation_times(context: Any) -> tuple[datetime, datetime]:
    """从可恢复状态计算本批次首次开始和最后结束时间。"""
    invocations = [
        item
        for item in context.state.snapshot().get("invocations", [])
        if item.get("command") != "report"
    ]
    started = [item.get("started_at") for item in invocations if item.get("started_at")]
    finished = [item.get("finished_at") for item in invocations if item.get("finished_at")]
    if context.command != "report":
        started.append(context.started_at.isoformat(timespec="seconds"))
        finished.append(datetime.now().isoformat(timespec="seconds"))
    try:
        return datetime.fromisoformat(min(started)), datetime.fromisoformat(max(finished))
    except (TypeError, ValueError):
        now = datetime.now()
        return now, now


def _pipeline_elapsed_seconds(context: Any) -> float:
    """累计已执行阶段耗时，排除人工检查和重启间隔。"""
    stages = tuple(context.state.stages().values())
    total = sum(stage.accumulated_seconds for stage in stages)
    for stage in stages:
        if stage.name != "report" or stage.status.value != "running" or not stage.started_at:
            continue
        try:
            total += max(0.0, (datetime.now() - datetime.fromisoformat(stage.started_at)).total_seconds())
        except ValueError:
            continue
    return total


def _matching_row(sheet: Any, values: list[Any], keys: tuple[int, ...]) -> int | None:
    """按业务键查找已有数据行，忽略仅有样式的空白行。"""
    for row in range(3, sheet.max_row + 1):
        if all(str(sheet.cell(row, key).value) == str(values[key - 1]) for key in keys):
            return row
    return None


def _last_business_row(sheet: Any) -> int:
    """返回首列存在真实批次 ID 的最后一行。"""
    return max(
        (
            row
            for row in range(3, sheet.max_row + 1)
            if sheet.cell(row, 1).value not in (None, "")
        ),
        default=2,
    )


def _copy_row_template(sheet: Any, source_row: int, target_row: int) -> None:
    """复制相邻业务行的样式和公式，保持 Excel 模板格式。"""
    from openpyxl.formula.translate import Translator

    for column in range(1, sheet.max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy.copy(source._style)
        if isinstance(source.value, str) and source.value.startswith("="):
            target.value = Translator(
                source.value, origin=source.coordinate
            ).translate_formula(target.coordinate)
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height


def _upsert_row(
    sheet: Any,
    values: list[Any],
    keys: tuple[int, ...],
    overwrite_formula_indexes: tuple[int, ...] = (),
) -> int:
    """更新匹配业务行，或紧邻真实数据末行追加一行。"""
    target = _matching_row(sheet, values, keys)
    if target is None:
        target = _last_business_row(sheet) + 1
        _copy_row_template(sheet, max(2, target - 1), target)
    for column, value in enumerate(values, 1):
        cell = sheet.cell(target, column)
        if (
            isinstance(cell.value, str)
            and cell.value.startswith("=")
            and column - 1 not in overwrite_formula_indexes
            and value is None
        ):
            continue
        cell.value = value
    return target


def _statistics(
    context: Any,
    registry: Any,
) -> tuple[dict[str, int], dict[str, dict[str, int]]]:
    """读取统计；旧状态缺少检测统计时按共享注册表重建。"""
    metrics = registry.metrics
    if "detect" not in metrics:
        registry_path = context.config.registry_dir / "datasets.yaml"
        payload = _load_registry(registry_path)
        metrics = _collect_registered_statistics(payload["yolo"])
        statistics_path = context.work_dir / "reports" / "dataset_statistics.json"
        atomic_write_text(
            statistics_path,
            json.dumps(metrics, ensure_ascii=False, indent=2),
        )
        context.logger.warning(
            "REPORT_STATISTICS_REBUILT | reason=legacy_register_metrics | "
            "registry=%s | statistics=%s",
            registry_path,
            statistics_path,
        )
    images = {split: int(metrics.get("images", {}).get(split, 0)) for split in SPLITS}
    detect = metrics.get("detect", {})
    labels = {
        split: {
            name: int(detect.get(split, {}).get(name, 0))
            for name in CLASS_NAMES.values()
        }
        for split in SPLITS
    }
    return images, labels


def _build_report_rows(context: Any) -> dict[str, list[list[Any]]]:
    """构建纯检测批次总览、模型汇总和类别明细数据行。"""
    registry = _stage(context, "register_data")
    training = _stage(context, "train_detect")
    exported = _stage(context, "export_detect")
    images, labels = _statistics(context, registry)
    start_time, end_time = _invocation_times(context)
    elapsed = _pipeline_elapsed_seconds(context)
    training_metrics = training.metrics.get("metrics", {})
    class_metrics = {
        str(item.get("Class", item.get("class_name", ""))): item
        for item in training.metrics.get("class_metrics", [])
    }
    overview = [[
        context.config.run_name,
        "; ".join(_path(context, path.resolve()) for path in context.config.output_dirs),
        start_time,
        end_time,
        elapsed / 86400.0,
        images["train"],
        images["val"],
        images["train"] + images["val"],
        "是" if training.metrics.get("force_train") else "否",
        _path(context, training.artifacts.get("best")),
        "",
        "",
        _path(context, context.log_path),
        "纯检测流程",
    ]]
    model = [[
        context.config.run_name,
        "YOLO11-P2-Detect",
        "检测",
        images["train"],
        images["val"],
        sum(labels["train"].values()),
        sum(labels["val"].values()),
        training.metrics.get("planned_epochs"),
        training.metrics.get("actual_epoch"),
        training.metrics.get("best_epoch"),
        (training.metrics.get("training_seconds") or training.accumulated_seconds) / 86400.0,
        training_metrics.get("fitness"),
        _metric(training_metrics, "metrics/precision(B)"),
        _metric(training_metrics, "metrics/recall(B)"),
        _metric(training_metrics, "metrics/mAP50(B)"),
        _metric(training_metrics, "metrics/mAP50-95(B)"),
        "", "", "", "", "", "", "", "", "",
        training.metrics.get("batch"),
        "x".join(map(str, training.metrics.get("image_size", []))),
        _path(context, training.artifacts.get("initial_weight")),
        _path(context, exported.artifacts.get("onnx")),
        _path(context, context.log_path),
    ]]
    details: list[list[Any]] = []
    for class_id, name in CLASS_NAMES.items():
        item = class_metrics.get(name, {})
        details.append([
            context.config.run_name,
            "YOLO11-P2-Detect",
            "检测",
            "验证",
            class_id,
            name,
            labels["val"][name],
            _metric(item, "Box-P", "box_precision"),
            _metric(item, "Box-R", "box_recall"),
            _metric(item, "mAP50", "Box-mAP50"),
            _metric(item, "mAP50-95", "Box-mAP50-95"),
            "", "", "", "", "", "", "", "",
        ])
    return {"批次总览": overview, "模型汇总": model, "类别明细": details}


def _write_excel(context: Any, rows: dict[str, list[list[Any]]]) -> dict[str, list[int]]:
    """以模板校验、写锁和原子替换方式幂等写入三张报表。"""
    from openpyxl import load_workbook

    target = context.config.excel_path.expanduser().resolve()
    validate_excel_template(target)
    lock_path = target.with_suffix(f"{target.suffix}.pipeline.lock")
    descriptor = _acquire_lock(lock_path)
    temporary = target.with_name(f".{target.stem}.{context.config.run_name}.tmp{target.suffix}")
    workbook = None
    try:
        workbook = load_workbook(target, data_only=False)
        overview_rows = [
            _upsert_row(workbook["批次总览"], row, (1,), overwrite_formula_indexes=(4,))
            for row in rows["批次总览"]
        ]
        model_rows = [
            _upsert_row(workbook["模型汇总"], row, (1, 2))
            for row in rows["模型汇总"]
        ]
        class_rows = [
            _upsert_row(workbook["类别明细"], row, (1, 2, 4, 5))
            for row in rows["类别明细"]
        ]
        _normalize_excel_durations(workbook)
        _normalize_excel_metrics(workbook)
        workbook.save(temporary)
        workbook.close()
        workbook = None
        validate_excel_template(temporary)
        os.replace(temporary, target)
        return {"overview": overview_rows, "model": model_rows, "class": class_rows}
    except (OSError, PermissionError) as exc:
        raise ExcelPendingError(f"中心Excel暂时无法保存:{target},原因:{exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()
        if temporary.exists():
            temporary.unlink()
        os.close(descriptor)
        lock_path.unlink(missing_ok=True)


def write_report(context: Any) -> StageOutcome:
    """生成完整 JSON 报告并以原 Pipeline 规则维护中心 Excel。"""
    rows = _build_report_rows(context)
    reports_dir = context.work_dir / "reports"
    report_path = reports_dir / "training_report.json"
    atomic_write_text(report_path, json.dumps(rows, ensure_ascii=False, indent=2, default=str))
    try:
        row_numbers = _write_excel(context, rows)
    except Exception as exc:
        atomic_write_text(
            reports_dir / "excel_pending.json",
            json.dumps(
                {"excel": str(context.config.excel_path), "error": str(exc), "rows": rows},
                ensure_ascii=False,
                indent=2,
                default=str,
            ),
        )
        if isinstance(exc, ExcelPendingError):
            raise
        raise ExcelPendingError(f"中心Excel写入失败:{exc}") from exc
    (reports_dir / "excel_pending.json").unlink(missing_ok=True)
    context.logger.info(
        "EXCEL_ROW_UPSERT | overview=%s | model=%s | class_rows=%s",
        row_numbers["overview"],
        row_numbers["model"],
        row_numbers["class"],
    )
    for row in rows["模型汇总"]:
        context.logger.info(
            "MODEL_RESULT | model=%s | best_epoch=%s | actual_epoch=%s | "
            "fitness=%s | box_mAP50=%s | box_mAP50-95=%s",
            row[1], row[9], row[8], row[11], row[14], row[15],
        )
    for row in rows["类别明细"]:
        context.logger.info(
            "CLASS_RESULT | class_id=%s | class=%s | labels=%s | box_P=%s | "
            "box_R=%s | box_mAP50=%s | box_mAP50-95=%s",
            row[4], row[5], row[6], row[7], row[8], row[9], row[10],
        )
    return StageOutcome(
        message="中心Excel三表已更新",
        artifacts={"report": str(report_path), "excel": str(context.config.excel_path)},
        metrics={"overview_rows": 1, "model_rows": 1, "class_rows": len(rows["类别明细"])},
        value=rows,
    )

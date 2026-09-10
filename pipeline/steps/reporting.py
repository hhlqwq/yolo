"""固定三 Sheet Excel 的结果整理、幂等写入和失败补写."""

from __future__ import annotations

import copy
import json
import os
import socket
import time
from collections.abc import Iterable
from datetime import datetime
from pathlib import Path
from typing import Any

from ..core.context import RunContext
from ..core.errors import ExcelPendingError, PipelineError
from ..core.io_utils import atomic_write_text
from ..core.stage_runner import StageOutcome
from ..core.state import StageStatus


OVERVIEW_HEADERS = [
    "批次ID", "数据版本/来源", "开始时间", "结束时间", "总耗时", "训练图片数", "验证图片数", "总图片数",
    "是否强制重训", "检测best.pt", "分割best.pt", "PIDNet best.pt", "日志文件", "备注",
]
MODEL_HEADERS = [
    "批次ID", "模型", "任务", "训练图片", "验证图片", "训练标签总数", "验证标签总数", "计划Epoch", "实际Epoch",
    "最佳Epoch", "训练时长", "Fitness", "全类P", "全类R", "Box mAP50", "Box mAP50-95", "Mask P", "Mask R",
    "Mask mAP50", "Mask mAP50-95", "mIoU(All)", "mIoU(FG)", "Pixel Acc", "Mean Acc", "Macro Dice", "Batch",
    "ImgSz", "初始权重", "ONNX", "日志文件",
]
CLASS_HEADERS = [
    "批次ID", "模型", "任务", "数据划分", "类别ID", "类别名", "标签数", "Box P", "Box R", "Box mAP50",
    "Box mAP50-95", "Mask P", "Mask R", "Mask mAP50", "Mask mAP50-95", "IoU", "Dice", "类别像素Acc(=R)", "备注",
]
SHEET_HEADERS = {"批次总览": OVERVIEW_HEADERS, "模型汇总": MODEL_HEADERS, "类别明细": CLASS_HEADERS}
HEADER_ROW = 2
DATA_START_ROW = 3
EXCEL_DURATION_FORMAT = "[h]:mm:ss"
METRIC_COLUMNS = {
    "模型汇总": tuple(range(12, 26)),
    "类别明细": tuple(range(8, 19)),
}


def _duration(seconds: float | None) -> str:
    """把秒数转换为小时可超过 24 的 HH:MM:SS 文本."""
    total = max(0, int(round(seconds or 0.0)))
    hours, remainder = divmod(total, 3600)
    minutes, second = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{second:02d}"


def _format_image_size(value: Any) -> str:
    """将新旧训练记录中的尺寸统一为 Excel 文本。"""
    if isinstance(value, (list, tuple)):
        return "x".join(map(str, value))
    return str(value or "")


def _mean_metric(items: Iterable[dict[str, Any]], key: str) -> float | None:
    """计算已有逐类别指标的算术平均值."""
    values: list[float] = []
    for item in items:
        value = _normalize_percent_text(item.get(key))
        try:
            values.append(float(value))
        except (TypeError, ValueError):
            continue
    return sum(values) / len(values) if values else None


def _metric(metrics: dict[str, Any], *keys: str) -> Any:
    """按候选名称读取指标,兼容不同 Ultralytics 小版本."""
    for key in keys:
        value = metrics.get(key)
        if value is not None:
            return value
    return None


def _normalize_percent_text(value: Any) -> Any:
    """仅把历史百分比文本转换为等价小数,数值型指标保持原始精度."""
    if isinstance(value, str):
        text = value.strip()
        if not text.endswith("%"):
            return value
        try:
            return float(text[:-1]) / 100.0
        except ValueError:
            return value
    return value


def _path(context: RunContext, value: Any) -> str:
    """把报表路径统一转换成客户可访问的 UNC 展示路径."""
    if value in (None, ""):
        return ""
    return context.mapper.map_path(str(value))


def _json_default(value: Any) -> str:
    """把报表中的原生时间值转换为稳定的 ISO JSON 字符串."""
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    raise TypeError(f"不支持写入报表JSON的类型:{type(value).__name__}")


def _stage(context: RunContext, name: str):
    """读取阶段并在缺失时给出明确错误."""
    stage = context.state.get_stage(name)
    if stage.status not in {StageStatus.SUCCEEDED, StageStatus.SKIPPED}:
        raise PipelineError(f"生成报告要求阶段成功,当前{name}={stage.status}")
    if not stage.artifacts and not stage.metrics:
        raise PipelineError(f"生成报告缺少阶段结果:{name}")
    return stage


def _invocation_times(context: RunContext) -> tuple[datetime, datetime, float]:
    """返回本批次第一次启动、最后结束和累计调用跨度."""
    invocations = [
        item
        for item in context.state.snapshot().get("invocations", [])
        if item.get("command") != "report"
    ]
    started = [item.get("started_at") for item in invocations if item.get("started_at")]
    finished = [item.get("finished_at") for item in invocations if item.get("finished_at")]
    if context.command != "report":
        started.append(context.started_at.isoformat(timespec="seconds"))
        finished.append(datetime.now().isoformat(timespec="seconds"))
    start_text = min(started) if started else datetime.now().isoformat(timespec="seconds")
    end_text = max(finished) if finished else datetime.now().isoformat(timespec="seconds")
    try:
        start_time = datetime.fromisoformat(start_text)
        end_time = datetime.fromisoformat(end_text)
        elapsed = (end_time - start_time).total_seconds()
    except ValueError:
        start_time = datetime.now()
        end_time = start_time
        elapsed = 0.0
    return start_time, end_time, elapsed


def _pipeline_elapsed_seconds(context: RunContext) -> float:
    """累计全部已执行阶段的耗时，排除失败后人工检查与重启间隔。"""
    stages = tuple(context.state.stages().values())
    total = sum(stage.accumulated_seconds for stage in stages)
    for stage in stages:
        if stage.name != "report" or stage.status != StageStatus.RUNNING or not stage.started_at:
            continue
        try:
            current_elapsed = (datetime.now() - datetime.fromisoformat(stage.started_at)).total_seconds()
        except ValueError:
            current_elapsed = 0.0
        total += max(0.0, current_elapsed)
    return total


def build_report_rows(context: RunContext) -> dict[str, list[list[Any]]]:
    """使用固定模板构建批次、模型和逐类别行."""
    registry = _stage(context, "register_data")
    statistics = registry.metrics
    detect = _stage(context, "train_detect")
    segment = _stage(context, "train_segment")
    pidnet = _stage(context, "train_pidnet")
    exports = {
        "detect": _stage(context, "export_detect"),
        "segment": _stage(context, "export_segment"),
        "pidnet": _stage(context, "export_pidnet"),
    }
    train_images = int(statistics["images"]["train"])
    val_images = int(statistics["images"]["val"])
    start_text, end_text, _ = _invocation_times(context)
    elapsed = _pipeline_elapsed_seconds(context)
    overview = [[
        context.config.run_name, "; ".join(_path(context, path.resolve()) for path in context.config.output_dirs), start_text, end_text,
        elapsed, train_images, val_images, train_images + val_images,
        "是" if any(stage.metrics.get("force_train") for stage in (detect, segment, pidnet)) else "否",
        _path(context, detect.artifacts.get("best")),
        _path(context, segment.artifacts.get("best")), _path(context, pidnet.artifacts.get("best")),
        _path(context, context.log_path), "",
    ]]
    model_rows: list[list[Any]] = []
    class_rows: list[list[Any]] = []
    for task, model_name, task_name, class_names, stage in (
        ("detect", "YOLO11-P2-Detect", "检测", {"paper": 0, "liquid": 1, "metal": 2}, detect),
        ("segment", "YOLO11-P2-Segment", "实例分割", {"paper": 0}, segment),
    ):
        payload = stage.metrics
        metrics = payload.get("metrics", {})
        counts = statistics[task]
        epochs = payload.get("planned_epochs")
        batch = payload.get("batch")
        duration = payload.get("training_seconds") or stage.accumulated_seconds
        model_rows.append([
            context.config.run_name, model_name, task_name, train_images, val_images,
            sum(counts["train"].values()), sum(counts["val"].values()), epochs,
            payload.get("actual_epoch"), payload.get("best_epoch"), duration, payload.get("fitness"),
            _metric(metrics, "metrics/precision(B)"), _metric(metrics, "metrics/recall(B)"),
            _metric(metrics, "metrics/mAP50(B)"), _metric(metrics, "metrics/mAP50-95(B)"),
            _metric(metrics, "metrics/precision(M)"), _metric(metrics, "metrics/recall(M)"),
            _metric(metrics, "metrics/mAP50(M)"), _metric(metrics, "metrics/mAP50-95(M)"),
            "", "", "", "", "", batch, _format_image_size(payload.get("image_size")),
            _path(context, stage.artifacts.get("initial_weight")), _path(context, exports[task].artifacts.get("onnx")),
            _path(context, context.log_path),
        ])
        items_by_name = {
            str(item.get("Class", item.get("class_name", ""))): item
            for item in payload.get("class_metrics", [])
        }
        for class_name, class_id in class_names.items():
            item = items_by_name.get(class_name, {})
            class_rows.append([
                context.config.run_name, model_name, task_name, "验证", class_id, class_name,
                counts["val"].get(class_name, 0),
                _metric(item, "Box-P", "box_precision"), _metric(item, "Box-R", "box_recall"),
                _metric(item, "mAP50", "Box-mAP50"), _metric(item, "mAP50-95", "Box-mAP50-95"),
                _metric(item, "Mask-P", "mask_precision"), _metric(item, "Mask-R", "mask_recall"),
                _metric(item, "Mask-mAP50"), _metric(item, "Mask-mAP50-95"), "", "", "", "",
            ])

    pid_payload = pidnet.metrics
    semantic = pid_payload.get("metrics", {})
    pid_items = {int(item.get("class_id", -1)): item for item in semantic.get("class_metrics", [])}
    foreground_items = [pid_items[class_id] for class_id in (1, 2) if class_id in pid_items]
    foreground_precision = _mean_metric(foreground_items, "precision")
    foreground_recall = _mean_metric(foreground_items, "recall")
    foreground_iou = _mean_metric(foreground_items, "iou")
    foreground_dice = _mean_metric(foreground_items, "dice")
    model_rows.append([
        context.config.run_name, "PIDNet", "语义分割", train_images, val_images, train_images, val_images,
        pid_payload.get("planned_epochs"), pid_payload.get("actual_epoch"), pid_payload.get("best_epoch"),
        pidnet.accumulated_seconds, "", foreground_precision, foreground_recall,
        "", "", "", "", "", "", "", foreground_iou,
        "", foreground_recall, foreground_dice,
        pid_payload.get("batch"), _format_image_size(pid_payload.get("image_size")),
        _path(context, pidnet.artifacts.get("initial_weight")), _path(context, exports["pidnet"].artifacts.get("onnx")),
        _path(context, context.log_path),
    ])
    pid_counts = statistics["pidnet_pixels"]["val"]
    for class_name, class_id in {"liquid": 1, "metal": 2}.items():
        item = pid_items.get(class_id, {})
        class_rows.append([
            context.config.run_name, "PIDNet", "语义分割", "验证", class_id, class_name,
            pid_counts.get(class_name, item.get("pixels", 0)), item.get("precision"), item.get("recall"),
            "", "", "", "",
            "", "", item.get("iou"), item.get("dice"), item.get("recall"), "类别像素准确率等同于该类别Recall.",
        ])
    return {"批次总览": overview, "模型汇总": model_rows, "类别明细": class_rows}


def _find_row(sheet: Any, key_indexes: tuple[int, ...], row: list[Any]) -> int | None:
    """根据业务主键查找已有数据行."""
    expected = tuple(str(row[index]) for index in key_indexes)
    for row_number in range(DATA_START_ROW, sheet.max_row + 1):
        actual = tuple(
            "" if sheet.cell(row_number, index + 1).value is None else str(sheet.cell(row_number, index + 1).value)
            for index in key_indexes
        )
        if actual == expected:
            return row_number
    return None


def _copy_row_template(sheet: Any, source_row: int, target_row: int) -> None:
    """复制上一行样式和公式,不改变模板结构."""
    if source_row < 2:
        return
    from openpyxl.formula.translate import Translator

    for column in range(1, sheet.max_column + 1):
        source, target = sheet.cell(source_row, column), sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy.copy(source._style)
        if isinstance(source.value, str) and source.value.startswith("="):
            target.value = Translator(source.value, origin=source.coordinate).translate_formula(target.coordinate)
    if source_row in sheet.row_dimensions:
        sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height


def _upsert_row(
    sheet: Any,
    key_indexes: tuple[int, ...],
    row: list[Any],
    overwrite_formula_indexes: tuple[int, ...] = (),
) -> None:
    """按主键更新或追加一行,保留模板公式列."""
    row_number = _find_row(sheet, key_indexes, row)
    if row_number is None:
        row_number = next(
            (
                candidate
                for candidate in range(DATA_START_ROW, sheet.max_row + 1)
                if all(sheet.cell(candidate, index + 1).value is None for index in key_indexes)
            ),
            sheet.max_row + 1,
        )
        if row_number > sheet.max_row:
            _copy_row_template(sheet, row_number - 1, row_number)
    for column, value in enumerate(row, start=1):
        cell = sheet.cell(row_number, column)
        if column - 1 in overwrite_formula_indexes or not (
            isinstance(cell.value, str) and cell.value.startswith("=")
        ):
            cell.value = value


def _excel_duration(seconds: Any) -> Any:
    """把秒数转换为 Excel 的天数值,空值保持不变."""
    if seconds in (None, ""):
        return seconds
    return max(0.0, float(seconds)) / 86400.0


def _parse_duration_text(value: Any) -> Any:
    """把历史 HH:MM:SS 文本转换为 Excel 天数值."""
    if not isinstance(value, str) or value.startswith("="):
        return value
    parts = value.strip().split(":")
    if len(parts) != 3:
        return value
    try:
        hours, minutes, seconds = (float(part) for part in parts)
    except ValueError:
        return value
    if hours < 0 or not 0 <= minutes < 60 or not 0 <= seconds < 60:
        return value
    return (hours * 3600 + minutes * 60 + seconds) / 86400.0


def _rows_for_excel(rows: dict[str, list[list[Any]]]) -> dict[str, list[list[Any]]]:
    """复制报表行并仅把 Excel 时长列转换为天数值."""
    result = copy.deepcopy(rows)
    for row in result["批次总览"]:
        row[4] = _excel_duration(row[4])
    for row in result["模型汇总"]:
        row[10] = _excel_duration(row[10])
    return result


def _normalize_excel_durations(workbook: Any) -> None:
    """迁移历史时长文本并应用支持超过 24 小时的显示格式."""
    for sheet_name, column in (("批次总览", 5), ("模型汇总", 11)):
        sheet = workbook[sheet_name]
        for row_number in range(DATA_START_ROW, sheet.max_row + 1):
            cell = sheet.cell(row_number, column)
            cell.value = _parse_duration_text(cell.value)
            cell.number_format = EXCEL_DURATION_FORMAT


def _foreground_metrics_from_sheet(sheet: Any) -> dict[str, dict[str, float | None]]:
    """从类别明细读取各批次 PIDNet 的 liquid/metal 前景平均指标."""
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row_number in range(DATA_START_ROW, sheet.max_row + 1):
        if sheet.cell(row_number, 2).value != "PIDNet":
            continue
        class_id = sheet.cell(row_number, 5).value
        if class_id not in (1, 2, "1", "2"):
            continue
        grouped.setdefault(str(sheet.cell(row_number, 1).value), []).append(
            {
                "precision": sheet.cell(row_number, 8).value,
                "recall": sheet.cell(row_number, 9).value,
                "iou": sheet.cell(row_number, 16).value,
                "dice": sheet.cell(row_number, 17).value,
            }
        )
    return {
        run_name: {
            "precision": _mean_metric(items, "precision"),
            "recall": _mean_metric(items, "recall"),
            "iou": _mean_metric(items, "iou"),
            "dice": _mean_metric(items, "dice"),
        }
        for run_name, items in grouped.items()
    }


def _remove_pidnet_background_rows(sheet: Any) -> None:
    """删除历史 PIDNet background 明细行并同步缩短表格区域."""
    rows = [
        row_number
        for row_number in range(DATA_START_ROW, sheet.max_row + 1)
        if sheet.cell(row_number, 2).value == "PIDNet"
        and (
            sheet.cell(row_number, 5).value in (0, "0")
            or str(sheet.cell(row_number, 6).value).lower() == "background"
        )
    ]
    for row_number in reversed(rows):
        sheet.delete_rows(row_number, 1)
    if rows:
        from openpyxl.utils.cell import get_column_letter, range_boundaries

        for table in sheet.tables.values():
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            removed = sum(min_row <= row_number <= max_row for row_number in rows)
            if removed:
                table.ref = (
                    f"{get_column_letter(min_col)}{min_row}:"
                    f"{get_column_letter(max_col)}{max_row - removed}"
                )


def _normalize_pidnet_excel_metrics(workbook: Any) -> None:
    """使新旧 PIDNet Excel 汇总均只统计 liquid/metal 前景类别."""
    class_sheet = workbook["类别明细"]
    _remove_pidnet_background_rows(class_sheet)
    foreground = _foreground_metrics_from_sheet(class_sheet)
    model_sheet = workbook["模型汇总"]
    for row_number in range(DATA_START_ROW, model_sheet.max_row + 1):
        if model_sheet.cell(row_number, 2).value != "PIDNet":
            continue
        metrics = foreground.get(str(model_sheet.cell(row_number, 1).value), {})
        model_sheet.cell(row_number, 13).value = metrics.get("precision")
        model_sheet.cell(row_number, 14).value = metrics.get("recall")
        model_sheet.cell(row_number, 21).value = None
        model_sheet.cell(row_number, 22).value = metrics.get("iou")
        model_sheet.cell(row_number, 23).value = None
        model_sheet.cell(row_number, 24).value = metrics.get("recall")
        model_sheet.cell(row_number, 25).value = metrics.get("dice")


def _normalize_excel_metrics(workbook: Any) -> None:
    """让中心 Excel 的全部新旧指标显示 3 位小数且不显示百分号."""
    for sheet_name, columns in METRIC_COLUMNS.items():
        sheet = workbook[sheet_name]
        for row_number in range(DATA_START_ROW, sheet.max_row + 1):
            for column in columns:
                cell = sheet.cell(row_number, column)
                if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                    cell.value = _normalize_percent_text(cell.value)
                cell.number_format = "0.000"


def _validate_headers(workbook: Any) -> None:
    """严格校验固定 Sheet 和表头,防止写错工作簿."""
    for sheet_name, expected in SHEET_HEADERS.items():
        if sheet_name not in workbook.sheetnames:
            raise PipelineError(f"Excel缺少固定Sheet:{sheet_name}")
        sheet = workbook[sheet_name]
        actual = [sheet.cell(HEADER_ROW, index).value for index in range(1, len(expected) + 1)]
        if actual != expected:
            raise PipelineError(f"Excel表头不匹配:{sheet_name},期望={expected},实际={actual}")


def validate_excel_template(path: Path) -> None:
    """只读打开并校验中心 Excel 模板."""
    if not path.is_file():
        raise PipelineError(f"中心Excel不存在:{path}")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise PipelineError("当前环境缺少openpyxl,无法维护中心Excel") from exc
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        _validate_headers(workbook)
    finally:
        workbook.close()


def _acquire_lock(lock_path: Path) -> int:
    """获取同目录排他锁,避免多个训练批次同时覆盖 Excel."""
    if lock_path.exists() and time.time() - lock_path.stat().st_mtime > 12 * 3600:
        lock_path.unlink()
    try:
        descriptor = os.open(lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    except FileExistsError as exc:
        raise ExcelPendingError(f"中心Excel正在被其他pipeline写入,锁文件:{lock_path}") from exc
    os.write(descriptor, f"host={socket.gethostname()},pid={os.getpid()}".encode("utf-8"))
    return descriptor


def write_excel(context: RunContext, rows: dict[str, list[list[Any]]]) -> None:
    """严格按固定模板幂等写入中心 Excel,并使用同目录原子替换."""
    from openpyxl import load_workbook

    target = context.config.excel_path.expanduser().resolve()
    validate_excel_template(target)
    lock_path = target.with_suffix(f"{target.suffix}.pipeline.lock")
    descriptor = _acquire_lock(lock_path)
    temporary = target.with_name(f".{target.stem}.{context.config.run_name}.tmp{target.suffix}")
    workbook = None
    try:
        workbook = load_workbook(target, data_only=False)
        _validate_headers(workbook)
        excel_rows = _rows_for_excel(rows)
        key_indexes = {"批次总览": (0,), "模型汇总": (0, 1), "类别明细": (0, 1, 3, 4)}
        for sheet_name, sheet_rows in excel_rows.items():
            for row in sheet_rows:
                overwrite_formula_indexes = (4,) if sheet_name == "批次总览" else ()
                _upsert_row(workbook[sheet_name], key_indexes[sheet_name], row, overwrite_formula_indexes)
        _normalize_pidnet_excel_metrics(workbook)
        _normalize_excel_durations(workbook)
        _normalize_excel_metrics(workbook)
        workbook.save(temporary)
        workbook.close()
        workbook = None
        validate_excel_template(temporary)
        os.replace(temporary, target)
    except (OSError, PermissionError) as exc:
        raise ExcelPendingError(f"中心Excel暂时无法保存:{target},原因:{exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()
        if temporary.exists():
            temporary.unlink()
        os.close(descriptor)
        lock_path.unlink(missing_ok=True)


def write_report(context: RunContext) -> StageOutcome:
    """生成报告快照并写入唯一中心 Excel;失败时保存待补写文件."""
    rows = build_report_rows(context)
    reports_dir = context.work_dir / "reports"
    reports_dir.mkdir(parents=True, exist_ok=True)
    snapshot_path = reports_dir / "training_report.json"
    pending_path = reports_dir / "excel_pending.json"
    atomic_write_text(snapshot_path, json.dumps(rows, ensure_ascii=False, indent=2, default=_json_default))
    for row in rows["模型汇总"]:
        context.logger.info(
            "MODEL_RESULT | model=%s | best_epoch=%s | actual_epoch=%s | fitness=%s | "
            "box_mAP50=%s | box_mAP50-95=%s | mask_mAP50=%s | mask_mAP50-95=%s | "
            "mIoU_all=%s | mIoU_fg=%s | pixel_acc=%s | duration=%s",
            row[1], row[9], row[8], row[11], row[14], row[15], row[18], row[19],
            row[20], row[21], row[22], _duration(row[10]),
        )
    for row in rows["类别明细"]:
        context.logger.info(
            "CLASS_RESULT | model=%s | class_id=%s | class=%s | labels=%s | box_P=%s | box_R=%s | "
            "box_mAP50=%s | box_mAP50-95=%s | mask_P=%s | mask_R=%s | mask_mAP50=%s | "
            "mask_mAP50-95=%s | IoU=%s | Dice=%s",
            row[1], row[4], row[5], row[6], row[7], row[8], row[9], row[10], row[11], row[12],
            row[13], row[14], row[15], row[16],
        )
    try:
        write_excel(context, rows)
    except Exception as exc:
        atomic_write_text(
            pending_path,
            json.dumps(
                {
                    "excel": str(context.config.excel_path),
                    "created_at": datetime.now().isoformat(timespec="seconds"),
                    "error": str(exc),
                    "rows": rows,
                },
                ensure_ascii=False,
                indent=2,
                default=_json_default,
            ),
        )
        if isinstance(exc, ExcelPendingError):
            raise
        raise ExcelPendingError(f"中心Excel写入失败:{exc}") from exc
    pending_path.unlink(missing_ok=True)
    return StageOutcome(
        message="中心Excel已更新",
        artifacts={"excel": str(context.config.excel_path), "report": str(snapshot_path)},
        metrics={"overview_rows": 1, "model_rows": len(rows["模型汇总"]), "class_rows": len(rows["类别明细"])},
        value=rows,
    )
